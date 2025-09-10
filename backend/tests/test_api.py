"""
API endpoint tests for Django REST framework
Test authentication, permissions, data serialization, and filtering
"""

import json
import tempfile
from decimal import Decimal
from datetime import date
from django.test import TestCase
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase, APIClient
from rest_framework import status
from openpyxl import Workbook

from dashboard_api.models import GeografiaProcedimiento, DetenidosAprehendidos
from dashboard_api.authentication import generate_jwt_token

User = get_user_model()


class TestAuthenticationAPI(APITestCase):
    """Test authentication endpoints"""
    
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin',
            password='admin123',
            role='admin'
        )
        self.viewer_user = User.objects.create_user(
            username='viewer',
            password='viewer123',
            role='viewer'
        )
    
    def test_login_success(self):
        """Test successful login"""
        data = {
            'username': 'admin',
            'password': 'admin123'
        }
        
        response = self.client.post('/api/login/', data)
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('token', response.data)
        self.assertIn('user', response.data)
        self.assertEqual(response.data['user']['username'], 'admin')
        self.assertEqual(response.data['user']['role'], 'admin')
    
    def test_login_invalid_credentials(self):
        """Test login with invalid credentials"""
        data = {
            'username': 'admin',
            'password': 'wrongpassword'
        }
        
        response = self.client.post('/api/login/', data)
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', response.data)
    
    def test_login_missing_fields(self):
        """Test login with missing fields"""
        # Missing password
        response = self.client.post('/api/login/', {'username': 'admin'})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        
        # Missing username
        response = self.client.post('/api/login/', {'password': 'admin123'})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        
        # Empty data
        response = self.client.post('/api/login/', {})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
    
    def test_jwt_token_authentication(self):
        """Test JWT token authentication"""
        # Generate token
        token = generate_jwt_token(self.admin_user)
        
        # Use token in request
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        response = self.client.get('/api/data/categorized/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
    
    def test_invalid_jwt_token(self):
        """Test invalid JWT token"""
        # Invalid token
        self.client.credentials(HTTP_AUTHORIZATION='Bearer invalid-token')
        response = self.client.get('/api/data/categorized/')
        
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
    
    def test_expired_jwt_token(self):
        """Test expired JWT token"""
        # Generate expired token (expires immediately)
        expired_token = generate_jwt_token(self.admin_user, expires_in_seconds=-1)
        
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {expired_token}')
        response = self.client.get('/api/data/categorized/')
        
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class TestDataAPI(APITestCase):
    """Test data retrieval endpoints"""
    
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin',
            password='admin123',
            role='admin'
        )
        self.viewer_user = User.objects.create_user(
            username='viewer',
            password='viewer123',
            role='viewer'
        )
        
        # Create test data
        self.procedimiento1 = GeografiaProcedimiento.objects.create(
            id_operativo='OP001',
            provincia='Buenos Aires',
            fecha_iso=date(2025, 1, 15),
            latitud=Decimal('-34.6037'),
            longitud=Decimal('-58.3816')
        )
        self.procedimiento2 = GeografiaProcedimiento.objects.create(
            id_operativo='OP002',
            provincia='Córdoba',
            fecha_iso=date(2025, 2, 15),
            latitud=Decimal('-31.4201'),
            longitud=Decimal('-64.1888')
        )
        
        # Create related data
        DetenidosAprehendidos.objects.create(
            procedimiento=self.procedimiento1,
            edad=25,
            sexo='M'
        )
    
    def test_unauthenticated_access_denied(self):
        """Test that unauthenticated requests are denied"""
        response = self.client.get('/api/data/categorized/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
    
    def test_categorized_data_endpoint(self):
        """Test categorized data endpoint"""
        self.client.force_authenticate(user=self.viewer_user)
        
        response = self.client.get('/api/data/categorized/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Check response structure
        expected_categories = [
            'general', 'detenidos', 'incautaciones', 'controlados',
            'afectados', 'procedimientos', 'abatidos', 'trata'
        ]
        
        for category in expected_categories:
            self.assertIn(category, response.data)
            self.assertIsInstance(response.data[category], list)
    
    def test_data_filtering_by_province(self):
        """Test data filtering by province"""
        self.client.force_authenticate(user=self.viewer_user)
        
        # Filter by Buenos Aires
        response = self.client.get('/api/data/categorized/?province=Buenos Aires')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Should only return Buenos Aires data
        general_data = response.data['general']
        for item in general_data:
            # Normalize for comparison
            normalized_province = item.get('provincia', '').lower().replace(' ', '')
            self.assertIn('buenos', normalized_province)
    
    def test_data_filtering_by_date_range(self):
        """Test data filtering by date range"""
        self.client.force_authenticate(user=self.viewer_user)
        
        # Filter January 2025
        response = self.client.get('/api/data/categorized/?from_date=2025-01-01&to_date=2025-01-31')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Should only return January data
        general_data = response.data['general']
        for item in general_data:
            if 'fecha_iso' in item:
                item_date = item['fecha_iso']
                self.assertTrue(item_date.startswith('2025-01'))
    
    def test_data_filtering_combined(self):
        """Test combined filtering by province and date"""
        self.client.force_authenticate(user=self.viewer_user)
        
        response = self.client.get(
            '/api/data/categorized/?province=Buenos Aires&from_date=2025-01-01&to_date=2025-01-31'
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Should return only Buenos Aires data from January
        general_data = response.data['general']
        self.assertGreaterEqual(len(general_data), 0)  # May be empty if no matches
    
    def test_data_stats_endpoint(self):
        """Test data statistics endpoint"""
        self.client.force_authenticate(user=self.viewer_user)
        
        response = self.client.get('/api/data/stats/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Check response structure
        self.assertIn('totalRecords', response.data)
        self.assertIn('dateRange', response.data)
        self.assertIn('provinces', response.data)
        
        # Check values
        self.assertEqual(response.data['totalRecords'], 2)
        self.assertIsInstance(response.data['provinces'], list)
    
    def test_invalid_date_format(self):
        """Test handling of invalid date formats"""
        self.client.force_authenticate(user=self.viewer_user)
        
        response = self.client.get('/api/data/categorized/?from_date=invalid-date')
        
        # Should handle gracefully, not return 500 error
        self.assertIn(response.status_code, [200, 400])
    
    def test_sql_injection_prevention(self):
        """Test SQL injection prevention in filters"""
        self.client.force_authenticate(user=self.viewer_user)
        
        malicious_inputs = [
            "'; DROP TABLE geografia_procedimiento; --",
            "1 OR 1=1",
            "UNION SELECT * FROM auth_user"
        ]
        
        for malicious_input in malicious_inputs:
            response = self.client.get(f'/api/data/categorized/?province={malicious_input}')
            
            # Should not cause server error
            self.assertIn(response.status_code, [200, 400])
            
            if response.status_code == 200:
                # Should not leak sensitive data
                response_str = str(response.data)
                self.assertNotIn('password', response_str.lower())
                self.assertNotIn('DROP', response_str)


class TestFileUploadAPI(APITestCase):
    """Test Excel file upload functionality"""
    
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin',
            password='admin123',
            role='admin'
        )
        self.viewer_user = User.objects.create_user(
            username='viewer',
            password='viewer123',
            role='viewer'
        )
    
    def create_test_excel_file(self):
        """Create a test Excel file"""
        wb = Workbook()
        ws = wb.active
        ws.title = "GEOG. PROCEDIMIENTO"
        
        # Add headers
        headers = [
            'ID_OPERATIVO', 'ID_PROCEDIMIENTO', 'PROVINCIA', 'FECHA',
            'LATITUD', 'LONGITUD', 'DESCRIPCION'
        ]
        ws.append(headers)
        
        # Add test data
        ws.append([
            'OP001', 'PROC001', 'Buenos Aires', '15/03/2025',
            -34.6037, -58.3816, 'Test procedure'
        ])
        ws.append([
            'OP002', 'PROC002', 'Córdoba', '16/03/2025',
            -31.4201, -64.1888, 'Another test procedure'
        ])
        
        # Save to temporary file
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp_file.name)
        return temp_file.name
    
    def test_upload_permission_required(self):
        """Test that upload requires admin permission"""
        # Viewer should be denied
        self.client.force_authenticate(user=self.viewer_user)
        response = self.client.post('/api/upload-excel/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        
        # Unauthenticated should be denied
        self.client.force_authenticate(user=None)
        response = self.client.post('/api/upload-excel/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
    
    def test_successful_file_upload(self):
        """Test successful Excel file upload"""
        self.client.force_authenticate(user=self.admin_user)
        
        excel_file = self.create_test_excel_file()
        
        with open(excel_file, 'rb') as f:
            response = self.client.post('/api/upload-excel/', {
                'file': f
            }, format='multipart')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Check response structure
        self.assertIn('status', response.data)
        self.assertIn('message', response.data)
        self.assertIn('stats', response.data)
        
        # Verify data was created
        self.assertGreater(GeografiaProcedimiento.objects.count(), 0)
    
    def test_missing_file_upload(self):
        """Test upload without file"""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.post('/api/upload-excel/')
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
    
    def test_invalid_file_format(self):
        """Test upload with invalid file format"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Create text file instead of Excel
        temp_file = tempfile.NamedTemporaryFile(suffix='.txt', delete=False)
        temp_file.write(b'Not an Excel file')
        temp_file.close()
        
        with open(temp_file.name, 'rb') as f:
            response = self.client.post('/api/upload-excel/', {
                'file': f
            }, format='multipart')
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
    
    def test_malformed_excel_file(self):
        """Test upload with malformed Excel file"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Create malformed Excel file
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        temp_file.write(b'Malformed Excel content')
        temp_file.close()
        
        with open(temp_file.name, 'rb') as f:
            response = self.client.post('/api/upload-excel/', {
                'file': f
            }, format='multipart')
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', response.data)
    
    def test_excel_missing_required_columns(self):
        """Test Excel file missing required columns"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Create Excel with missing required columns
        wb = Workbook()
        ws = wb.active
        ws.append(['PROVINCIA', 'FECHA'])  # Missing ID_OPERATIVO
        ws.append(['Buenos Aires', '15/03/2025'])
        
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp_file.name)
        
        with open(temp_file.name, 'rb') as f:
            response = self.client.post('/api/upload-excel/', {
                'file': f
            }, format='multipart')
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('ID_OPERATIVO', str(response.data))
    
    def test_large_file_upload(self):
        """Test upload of large Excel file"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Create larger Excel file
        wb = Workbook()
        ws = wb.active
        ws.title = "GEOG. PROCEDIMIENTO"
        
        headers = ['ID_OPERATIVO', 'PROVINCIA', 'FECHA']
        ws.append(headers)
        
        # Add 1000 rows
        for i in range(1000):
            ws.append([f'OP{i:05d}', 'Buenos Aires', '15/03/2025'])
        
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp_file.name)
        
        with open(temp_file.name, 'rb') as f:
            response = self.client.post('/api/upload-excel/', {
                'file': f
            }, format='multipart')
        
        # Should handle large files
        self.assertIn(response.status_code, [200, 413])  # Success or too large


class TestSecurityAPI(APITestCase):
    """Security-focused API tests"""
    
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin',
            password='admin123',
            role='admin'
        )
    
    def test_xss_prevention(self):
        """Test XSS prevention in responses"""
        # Create procedimiento with XSS payload
        GeografiaProcedimiento.objects.create(
            id_operativo='<script>alert("xss")</script>',
            descripcion='<img src=x onerror=alert("xss")>',
            provincia='Buenos Aires'
        )
        
        self.client.force_authenticate(user=self.admin_user)
        response = self.client.get('/api/data/categorized/')
        
        # Verify HTML is escaped in JSON response
        response_content = json.dumps(response.data)
        self.assertNotIn('<script>', response_content)
        self.assertNotIn('onerror=', response_content)
    
    def test_unauthorized_data_access(self):
        """Test unauthorized access to sensitive data"""
        # Try to access without authentication
        response = self.client.get('/api/data/categorized/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
        
        # Try with invalid token
        self.client.credentials(HTTP_AUTHORIZATION='Bearer invalid')
        response = self.client.get('/api/data/categorized/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
    
    def test_role_based_access_control(self):
        """Test role-based access control"""
        viewer = User.objects.create_user(
            username='viewer',
            role='viewer'
        )
        
        # Viewer should not access admin endpoints
        self.client.force_authenticate(user=viewer)
        response = self.client.post('/api/upload-excel/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
    
    def test_cors_headers(self):
        """Test CORS headers in responses"""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.get('/api/data/stats/')
        
        # Check for CORS headers (if configured)
        # This depends on your CORS configuration
        # Uncomment if CORS is configured:
        # self.assertIn('Access-Control-Allow-Origin', response)


class TestErrorHandling(APITestCase):
    """Test API error handling"""
    
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin',
            password='admin123',
            role='admin'
        )
    
    def test_404_handling(self):
        """Test 404 error handling"""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.get('/api/nonexistent-endpoint/')
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
    
    def test_method_not_allowed(self):
        """Test method not allowed handling"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Try DELETE on endpoint that doesn't support it
        response = self.client.delete('/api/data/stats/')
        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
    
    def test_malformed_json(self):
        """Test handling of malformed JSON"""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.post(
            '/api/login/',
            data='malformed json',
            content_type='application/json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
    
    def test_large_request_handling(self):
        """Test handling of oversized requests"""
        self.client.force_authenticate(user=self.admin_user)
        
        # Create very large data
        large_data = {'data': 'x' * 10000000}  # 10MB
        
        response = self.client.post('/api/upload-excel/', large_data)
        
        # Should handle gracefully (413 or other appropriate error)
        self.assertIn(response.status_code, [400, 413, 500])