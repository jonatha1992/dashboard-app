#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script para poblar las tablas dimensionales desde el archivo Datos_Dimensiones.xlsx
"""
import os
import sys
import django
import openpyxl
from datetime import datetime

# Configurar Django
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'dashboard_project.settings')
django.setup()

from dashboard_api.models import Regional, UnidadOperativa, TipoOperativo, TipoDelito


def poblar_regionales_y_unidades():
    """
    Poblar tablas Regional y UnidadOperativa desde la hoja UNIDADES
    """
    print("[INFO] Poblando Regionales y Unidades Operativas...")
    
    # Ruta al archivo Excel
    excel_path = os.path.join(os.path.dirname(__file__), '..', 'notebooks', 'Datos_Dimensiones.xlsx')
    
    if not os.path.exists(excel_path):
        print(f"❌ No se encontró el archivo: {excel_path}")
        return
    
    workbook = openpyxl.load_workbook(excel_path, read_only=True)
    
    # Leer hoja UNIDADES
    if 'UNIDADES' not in workbook.sheetnames:
        print("❌ No se encontró la hoja UNIDADES")
        return
    
    sheet = workbook['UNIDADES']
    regionales_created = {}
    unidades_created = 0
    
    # Iterar por las filas (saltando la cabecera)
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[1]:  # Si no hay código regional, saltar
            continue
            
        codigo_regional = row[1]  # Columna B: REGIONAL
        codigo_unidad = row[2]    # Columna C: UNIDAD
        provincia = row[3]        # Columna D: PROVINCIA
        departamento = row[4]     # Columna E: DEPARTAMENTO/PARTIDO/COMUNA
        
        # Crear Regional si no existe
        if codigo_regional not in regionales_created:
            regional, created = Regional.objects.get_or_create(
                codigo_regional=codigo_regional,
                defaults={
                    'nombre_regional': f'Regional {codigo_regional}',
                    'descripcion': f'Regional {codigo_regional} con cobertura en múltiples provincias'
                }
            )
            regionales_created[codigo_regional] = regional
            if created:
                print(f"  [OK] Regional creada: {regional}")
        else:
            regional = regionales_created[codigo_regional]
        
        # Crear Unidad Operativa
        unidad, created = UnidadOperativa.objects.get_or_create(
            codigo_unidad=codigo_unidad,
            defaults={
                'regional': regional,
                'nombre_unidad': f'Unidad {codigo_unidad}',
                'provincia': provincia or 'Sin especificar',
                'departamento_partido_comuna': departamento or 'Sin especificar'
            }
        )
        
        if created:
            unidades_created += 1
            print(f"  ✅ Unidad creada: {unidad}")
    
    print(f"📈 Resumen: {len(regionales_created)} regionales, {unidades_created} unidades operativas creadas")
    workbook.close()


def poblar_tipos_operativo():
    """
    Poblar tabla TipoOperativo desde la hoja CODIGO_OPERATIVO
    """
    print("📊 Poblando Tipos de Operativo...")
    
    excel_path = os.path.join(os.path.dirname(__file__), '..', 'notebooks', 'Datos_Dimensiones.xlsx')
    
    if not os.path.exists(excel_path):
        print(f"❌ No se encontró el archivo: {excel_path}")
        return
    
    workbook = openpyxl.load_workbook(excel_path, read_only=True)
    
    # Leer hoja CODIGO_OPERATIVO
    if 'CODIGO_OPERATIVO' not in workbook.sheetnames:
        print("❌ No se encontró la hoja CODIGO_OPERATIVO")
        return
    
    sheet = workbook['CODIGO_OPERATIVO']
    operativos_created = 0
    
    # Iterar por las filas (saltando la cabecera)
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si no hay código operativo, saltar
            continue
            
        codigo_operativo = int(row[0])  # Columna A: CODIGO_OPERATIVO
        nombre_operativo = row[1]       # Columna B: NOMBRE_OPERATIVO  
        alcance = row[2]                # Columna C: ALCANCE
        
        # Determinar categoría basada en el código
        categoria = None
        if 6000 <= codigo_operativo <= 6999:
            categoria = 'PFA'
        elif 3000 <= codigo_operativo <= 3999:
            categoria = 'GNA'
        elif 2000 <= codigo_operativo <= 2999:
            categoria = 'PSA'
        else:
            categoria = 'GENERAL'
        
        # Crear Tipo Operativo
        tipo_op, created = TipoOperativo.objects.get_or_create(
            codigo_operativo=codigo_operativo,
            defaults={
                'nombre_operativo': nombre_operativo or f'Operativo {codigo_operativo}',
                'alcance': alcance or 'Sin descripción disponible',
                'categoria': categoria
            }
        )
        
        if created:
            operativos_created += 1
            print(f"  ✅ Tipo operativo creado: {tipo_op}")
    
    print(f"📈 Resumen: {operativos_created} tipos de operativo creados")
    workbook.close()


def poblar_tipos_delito():
    """
    Poblar tabla TipoDelito con categorías predefinidas
    """
    print("📊 Poblando Tipos de Delito...")
    
    tipos_delito = [
        # Detenciones/Aprehensiones
        {'categoria': 'detencion', 'subcategoria': 'Narcotráfico', 'descripcion': 'Delitos relacionados con drogas', 'gravedad': 'grave'},
        {'categoria': 'detencion', 'subcategoria': 'Contrabando', 'descripcion': 'Introducción ilegal de mercaderías', 'gravedad': 'moderada'},
        {'categoria': 'detencion', 'subcategoria': 'Tráfico de armas', 'descripcion': 'Comercio ilegal de armamento', 'gravedad': 'muy_grave'},
        {'categoria': 'detencion', 'subcategoria': 'Documentos falsos', 'descripcion': 'Falsificación de documentos', 'gravedad': 'moderada'},
        
        # Incautaciones
        {'categoria': 'incautacion', 'subcategoria': 'Drogas', 'descripcion': 'Sustancias estupefacientes', 'gravedad': 'grave'},
        {'categoria': 'incautacion', 'subcategoria': 'Armas de fuego', 'descripcion': 'Armas y municiones', 'gravedad': 'muy_grave'},
        {'categoria': 'incautacion', 'subcategoria': 'Mercadería', 'descripcion': 'Productos de contrabando', 'gravedad': 'moderada'},
        {'categoria': 'incautacion', 'subcategoria': 'Vehículos', 'descripcion': 'Vehículos relacionados con delitos', 'gravedad': 'moderada'},
        {'categoria': 'incautacion', 'subcategoria': 'Dinero', 'descripcion': 'Efectivo de origen ilícito', 'gravedad': 'grave'},
        
        # Trata/Tráfico
        {'categoria': 'trata', 'subcategoria': 'Trata de personas', 'descripcion': 'Explotación de personas', 'gravedad': 'muy_grave'},
        {'categoria': 'trata', 'subcategoria': 'Tráfico de personas', 'descripcion': 'Traslado ilegal de personas', 'gravedad': 'grave'},
        {'categoria': 'trata', 'subcategoria': 'Explotación laboral', 'descripcion': 'Trabajo forzoso', 'gravedad': 'grave'},
        {'categoria': 'trata', 'subcategoria': 'Explotación sexual', 'descripcion': 'Explotación con fines sexuales', 'gravedad': 'muy_grave'},
        
        # Abatidos
        {'categoria': 'abatido', 'subcategoria': 'Enfrentamiento armado', 'descripcion': 'Persona abatida en enfrentamiento', 'gravedad': 'muy_grave'},
        {'categoria': 'abatido', 'subcategoria': 'Resistencia al arresto', 'descripcion': 'Resistencia violenta durante detención', 'gravedad': 'muy_grave'},
        
        # Otros delitos
        {'categoria': 'otros', 'subcategoria': 'Delitos ambientales', 'descripcion': 'Daños al medio ambiente', 'gravedad': 'moderada'},
        {'categoria': 'otros', 'subcategoria': 'Ciberdelitos', 'descripcion': 'Delitos informáticos', 'gravedad': 'moderada'},
        {'categoria': 'otros', 'subcategoria': 'Lavado de dinero', 'descripcion': 'Blanqueo de capitales', 'gravedad': 'grave'},
        
        # Otros eventos
        {'categoria': 'evento', 'subcategoria': 'Accidente de tránsito', 'descripcion': 'Siniestros viales', 'gravedad': 'leve'},
        {'categoria': 'evento', 'subcategoria': 'Rescate', 'descripcion': 'Operaciones de rescate', 'gravedad': 'leve'},
        {'categoria': 'evento', 'subcategoria': 'Emergencia médica', 'descripcion': 'Asistencia médica', 'gravedad': 'leve'},
    ]
    
    delitos_created = 0
    for tipo_data in tipos_delito:
        tipo_delito, created = TipoDelito.objects.get_or_create(
            categoria=tipo_data['categoria'],
            subcategoria=tipo_data['subcategoria'],
            defaults={
                'descripcion': tipo_data['descripcion'],
                'gravedad': tipo_data['gravedad']
            }
        )
        
        if created:
            delitos_created += 1
            print(f"  ✅ Tipo delito creado: {tipo_delito}")
    
    print(f"📈 Resumen: {delitos_created} tipos de delito creados")


def main():
    """
    Función principal para poblar todas las tablas dimensionales
    """
    print("🚀 Iniciando población de tablas dimensionales...")
    print("=" * 60)
    
    try:
        poblar_regionales_y_unidades()
        print("\n" + "-" * 60)
        
        poblar_tipos_operativo()
        print("\n" + "-" * 60)
        
        poblar_tipos_delito()
        print("\n" + "=" * 60)
        
        print("✅ Población completada exitosamente!")
        
        # Mostrar resumen
        print("\n📊 RESUMEN FINAL:")
        print(f"   🏢 Regionales: {Regional.objects.count()}")
        print(f"   🏛️  Unidades: {UnidadOperativa.objects.count()}")
        print(f"   🎯 Tipos Operativo: {TipoOperativo.objects.count()}")
        print(f"   ⚖️  Tipos Delito: {TipoDelito.objects.count()}")
        
    except Exception as e:
        print(f"❌ Error durante la población: {str(e)}")
        import traceback
        traceback.print_exc()


if __name__ == '__main__':
    main()