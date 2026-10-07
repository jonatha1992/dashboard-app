import logging
from decimal import Decimal, InvalidOperation
import openpyxl
from drf_spectacular.utils import extend_schema
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import (
    Abatidos,
    CodigoOperativo,
    DetenidosAprehendidos,
    Fallecidos,
    GeografiaProcedimiento,
    Incautaciones,
    OtrosDelitos,
    OtrosEventos,
    PersonalElementosAfectados,
    TrataTraficPersonas,
    VehiculosPersonasControladas,
)
from .serializers import FileUploadSerializer, UploadResultSerializer

logger = logging.getLogger(__name__)


@extend_schema(
    tags=['Gestión de Datos'],
    summary='Subir datos operacionales',
    description='Subir archivo Excel con datos operacionales. Solo administradores. '
                'Procesa automáticamente y distribuye a tablas especializadas.',
    request=FileUploadSerializer,
    responses={
        200: UploadResultSerializer,
        400: {
            'description': 'Archivo requerido o inválido',
            'example': {'error': 'Archivo requerido'}
        },
        403: {
            'description': 'Permisos insuficientes',
            'example': {'error': 'Permisos insuficientes'}
        },
        500: {
            'description': 'Error procesando archivo',
            'example': {'error': 'Error al procesar archivo: [detalle]'}
        }
    }
)
class DataUploadView(APIView):
    """
    Upload operational data (equivalent to Node.js '/api/data/upload')
    Admin only - CON DISTRIBUCION INTELIGENTE A TABLAS ESPECIALIZADAS
    """
    permission_classes = [permissions.IsAuthenticated]
    
    def post(self, request):
        if request.user.role != 'admin':
            return Response({
                'error': 'Permisos insuficientes'
            }, status=status.HTTP_403_FORBIDDEN)
        
        serializer = FileUploadSerializer(data=request.data)
        if not serializer.is_valid():
            return Response({
                'error': 'Archivo requerido'
            }, status=status.HTTP_400_BAD_REQUEST)
        
        file = serializer.validated_data['file']
        
        try:
            workbook = openpyxl.load_workbook(file, data_only=True)
            
            stats = {
                'totalAdded': 0,
                'duplicatesSkipped': 0,
                'sheetsProcessed': [],
                'specialized_tables': {
                    'geografia_procedimientos': 0,
                    'vehiculos_controladas': 0,
                    'personal_afectados': 0,
                    'detenidos': 0,
                    'incautaciones': 0,
                    'trata_personas': 0,
                    'otros_delitos': 0,
                    'otros_eventos': 0,
                    'fallecidos': 0,
                    'abatidos': 0,
                    'codigos_operativos': 0
                }
            }
            
            master_table_processed = False
            master_sheet_data = {}
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                headers = []
                sheet_data = []
                
                for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if row_num == 1:
                        headers = [str(cell) if cell is not None else f'Column_{i}' for i, cell in enumerate(row)]
                        continue
                    
                    if all(cell in (None, '') for cell in row):
                        continue
                    
                    row_dict = {}
                    for i, cell in enumerate(row):
                        if i < len(headers):
                            row_dict[headers[i]] = str(cell) if cell is not None else ''
                    
                    sheet_data.append(row_dict)
                
                if sheet_data:
                    master_sheet_data[sheet_name] = sheet_data
            
            for sheet_name, sheet_data in master_sheet_data.items():
                if 'GEOG' in sheet_name.upper() and 'PROCEDIMIENTO' in sheet_name.upper():
                    logger.info("Procesando tabla maestra: %s", sheet_name)
                    sheet_stats = self._process_sheet_data_to_specialized_tables(
                        sheet_data, sheet_name, file.name, stats
                    )
                    
                    stats['sheetsProcessed'].append({
                        'name': sheet_name,
                        'totalRows': len(sheet_data),
                        'added': sheet_stats['added'],
                        'skipped': sheet_stats['skipped']
                    })
                    
                    stats['totalAdded'] += sheet_stats['added']
                    stats['duplicatesSkipped'] += sheet_stats['skipped']
                    master_table_processed = True
                    break
            
            if not master_table_processed:
                logger.warning("No se encontró la hoja GEOG. PROCEDIMIENTO - procesando sin tabla maestra")
            
            for sheet_name, sheet_data in master_sheet_data.items():
                if 'GEOG' in sheet_name.upper() and 'PROCEDIMIENTO' in sheet_name.upper():
                    continue
                
                logger.info("Procesando hoja especializada: %s", sheet_name)
                sheet_stats = self._process_sheet_data_to_specialized_tables(
                    sheet_data, sheet_name, file.name, stats
                )
                
                stats['sheetsProcessed'].append({
                    'name': sheet_name,
                    'totalRows': len(sheet_data),
                    'added': sheet_stats['added'],
                    'skipped': sheet_stats['skipped']
                })
                
                stats['totalAdded'] += sheet_stats['added']
                stats['duplicatesSkipped'] += sheet_stats['skipped']
            
            etl_result = {
                'status': 'disabled', 
                'message': 'ETL system removed', 
                'auto_executed': False
            }
            
            filtering_details = self._generate_filtering_statistics(stats)
            
            return Response({
                'success': True,
                'message': 'Datos cargados exitosamente con nueva estructura',
                'stats': stats,
                'filtering_details': filtering_details,
                'total_records_processed': stats.get('totalAdded', 0),
                'total_records_created': sum(stats.get('specialized_tables', {}).values()),
                'etl_result': etl_result
            })
            
        except Exception as e:
            logger.exception("Error al procesar archivo Excel: %s", e)
            return Response({
                'error': f'Error al procesar archivo: {str(e)}'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
    
    def _process_sheet_data_to_specialized_tables(self, sheet_data, sheet_name, original_filename, global_stats):
        sheet_stats = {'added': 0, 'skipped': 0}
        sheet_upper = sheet_name.upper()
        is_master_sheet = 'GEOG' in sheet_upper and 'PROCEDIMIENTO' in sheet_upper
        
        for row_data in sheet_data:
            try:
                if is_master_sheet:
                    procedimiento = self._create_or_update_geografia_procedimiento(
                        row_data, sheet_name, original_filename
                    )
                    
                    if procedimiento:
                        sheet_stats['added'] += 1
                        global_stats['specialized_tables']['geografia_procedimientos'] += 1
                    else:
                        sheet_stats['skipped'] += 1
                        continue
                else:
                    id_operativo = self._safe_str(row_data.get('ID_OPERATIVO'))
                    id_procedimiento = self._safe_str(row_data.get('ID_PROCEDIMIENTO'))
                    
                    if not id_operativo or not id_procedimiento:
                        logger.warning("Registro en %s sin ID_OPERATIVO/ID_PROCEDIMIENTO válido - omitiendo", sheet_name)
                        sheet_stats['skipped'] += 1
                        continue
                    
                    procedimiento = GeografiaProcedimiento.objects.filter(
                        id_operativo=id_operativo,
                        id_procedimiento=id_procedimiento
                    ).first()
                    
                    if not procedimiento:
                        logger.warning("No se encontró procedimiento maestro para %s_%s en %s - omitiendo", id_operativo, id_procedimiento, sheet_name)
                        sheet_stats['skipped'] += 1
                        continue
                    
                    specialized_created = self._distribute_to_specialized_table(
                        procedimiento, row_data, sheet_name, global_stats
                    )
                    
                    if specialized_created:
                        sheet_stats['added'] += 1
                    else:
                        sheet_stats['skipped'] += 1
                        
            except Exception as e:
                logger.error("Error procesando fila en %s: %s", sheet_name, e)
                sheet_stats['skipped'] += 1
        
        return sheet_stats
    
    def _create_or_update_geografia_procedimiento(self, row_data, sheet_name, original_filename):
        field_mapping = {
            'FUERZA_INTERVINIENTE': 'fuerza_interviniente',
            'ID_OPERATIVO': 'id_operativo',
            'ID_PROCEDIMIENTO': 'id_procedimiento',
            'UNIDAD_INTERVINIENTE': 'unidad_interviniente',
            'DESCRIPCIÓN': 'descripcion',
            'TIPO_INTERVENCION': 'tipo_intervencion',
            'PROVINCIA': 'provincia',
            'DEPARTAMENTO O PARTIDO': 'departamento_o_partido',
            'LOCALIDAD': 'localidad',
            'DIRECCION': 'direccion',
            'ZONA_SEGURIDAD_FRONTERAS': 'zona_seguridad_fronteras',
            'PASO_FRONTERIZO': 'paso_fronterizo',
            'LATITUD': 'latitud',
            'LONGITUD': 'longitud',
            'FECHA': 'fecha',
            'HORA': 'hora',
            'OTRAS AGENCIAS INTERVINIENTES': 'otras_agencias_intervinientes',
            'Observaciones - Detalles': 'observaciones',
        }
        
        id_operativo = self._safe_str(row_data.get('ID_OPERATIVO'))
        id_procedimiento = self._safe_str(row_data.get('ID_PROCEDIMIENTO'))
        
        if not id_operativo or not id_procedimiento:
            return None
        
        record_key = f"{id_operativo}_{id_procedimiento}_{sheet_name}"
        existing = GeografiaProcedimiento.objects.filter(record_key=record_key).first()
        
        if existing:
            return existing
        
        procedimiento_data = {
            'hoja': sheet_name,
            'archivo_original': original_filename,
            'record_key': record_key
        }
        
        for excel_field, model_field in field_mapping.items():
            if excel_field in row_data:
                value = row_data[excel_field]
                if value and value != '-':
                    procedimiento_data[model_field] = self._safe_str(value)
        
        if 'LATITUD' in row_data:
            procedimiento_data['latitud'] = self._safe_decimal(row_data['LATITUD'])
        if 'LONGITUD' in row_data:
            procedimiento_data['longitud'] = self._safe_decimal(row_data['LONGITUD'])
        
        procedimiento = GeografiaProcedimiento.objects.create(**procedimiento_data)
        return procedimiento
    
    def _distribute_to_specialized_table(self, procedimiento, row_data, sheet_name, global_stats):
        sheet_upper = sheet_name.upper()
        
        try:
            if 'VEHI' in sheet_upper and 'PERSO' in sheet_upper and 'CONTROLADAS' in sheet_upper:
                return self._create_vehiculos_personas_controladas(procedimiento, row_data, global_stats)
            
            elif 'PERSONAL' in sheet_upper and 'ELEMENTOS' in sheet_upper and 'AFECTADOS' in sheet_upper:
                return self._create_personal_elementos_afectados(procedimiento, row_data, global_stats)
            
            elif 'DETENIDOS' in sheet_upper or 'APREHENDIDOS' in sheet_upper:
                return self._create_detenidos_aprehendidos(procedimiento, row_data, global_stats)
            
            elif 'INCAUTACIONES' in sheet_upper:
                return self._create_incautaciones(procedimiento, row_data, global_stats)
            
            elif 'TRATA' in sheet_upper or 'TRAFIC' in sheet_upper:
                return self._create_trata_trafico_personas(procedimiento, row_data, global_stats)
            
            elif 'OTROS DELITOS' in sheet_upper:
                return self._create_otros_delitos(procedimiento, row_data, global_stats)
            
            elif 'OTROS EVENTOS' in sheet_upper:
                return self._create_otros_eventos(procedimiento, row_data, global_stats)
            
            elif 'FALLECIDOS' in sheet_upper:
                return self._create_fallecidos(procedimiento, row_data, global_stats)
            
            elif 'ABATIDOS' in sheet_upper:
                return self._create_abatidos(procedimiento, row_data, global_stats)
            
            elif 'CODIGO' in sheet_upper and 'OPERATIVO' in sheet_upper:
                return self._create_codigo_operativo(procedimiento, row_data, global_stats)
            
            global_stats['specialized_tables']['geografia_procedimientos'] += 1
            return True
            
        except Exception as e:
            logger.exception("Error creando datos especializados: %s", e)
            return False
    
    def _create_vehiculos_personas_controladas(self, procedimiento, row_data, global_stats):
        vehiculos_controlados = row_data.get('VEHICULOS_CONTROLADOS')
        personas_controladas = row_data.get('PERSONAS_CONTROLADAS')
        
        if (vehiculos_controlados and str(vehiculos_controlados).strip() not in ['-', '']) or \
           (personas_controladas and str(personas_controladas).strip() not in ['-', '']):
            
            VehiculosPersonasControladas.objects.create(
                procedimiento=procedimiento,
                vehiculos_controlados=self._safe_int(vehiculos_controlados),
                personas_controladas=self._safe_int(personas_controladas),
                cant_averiguaciones_secuestro=self._safe_int(row_data.get('CANT_AVERIGUACIONES_SECUESTRO')),
                cant_solicitudes_antecedentes=self._safe_int(row_data.get('CANT_SOLICITUDES_ANTECEDENTES')),
                cant_embarcaciones_controladas=self._safe_int(row_data.get('CANT_EMBARCACIONES_CONTROLADAS'))
            )
            global_stats['specialized_tables']['vehiculos_controladas'] += 1
            return True
        return False
    
    def _create_personal_elementos_afectados(self, procedimiento, row_data, global_stats):
        cant_efectivos = self._safe_int(row_data.get('CANT_EFECTIVOS'))
        
        if cant_efectivos and cant_efectivos > 0:
            PersonalElementosAfectados.objects.create(
                procedimiento=procedimiento,
                cant_efectivos=cant_efectivos,
                cant_autos_camionetas=self._safe_int(row_data.get('CANT_AUTOS_CAMIONETAS')),
                cant_scanners=self._safe_int(row_data.get('CANT_SCANNERS')),
                cant_embarcaciones=self._safe_int(row_data.get('CANT_EMBARCACIONES')),
                cant_motos=self._safe_int(row_data.get('CANT_MOTOS')),
                cant_caballos=self._safe_int(row_data.get('CANT_CABALLOS')),
                cant_canes=self._safe_int(row_data.get('CANT_CANES')),
                cant_morphrapid=self._safe_int(row_data.get('CANT_MORPHRAPID')),
                cant_lpr=self._safe_int(row_data.get('CANT_LPR'))
            )
            global_stats['specialized_tables']['personal_afectados'] += 1
            return True
        return False
    
    def _create_detenidos_aprehendidos(self, procedimiento, row_data, global_stats):
        edad_raw = row_data.get('EDAD')
        
        if edad_raw and str(edad_raw).strip() not in ['-', '']:
            DetenidosAprehendidos.objects.create(
                procedimiento=procedimiento,
                edad=self._safe_int(edad_raw),
                sexo=self._safe_str(row_data.get('SEXO')),
                nacionalidad=self._safe_str(row_data.get('NACIONALIDAD')),
                situacion_procesal=self._safe_str(row_data.get('SITUACION_PROCESAL')),
                delito_imputado=self._safe_str(row_data.get('DELITO_IMPUTADO')),
                juzgado_interviniente=self._safe_str(row_data.get('JUZGADO_INTERVINIENTE')),
                caratula_causa=self._safe_str(row_data.get('CARATULA_CAUSA')),
                num_causa=self._safe_str(row_data.get('NUM_CAUSA'))
            )
            global_stats['specialized_tables']['detenidos'] += 1
            return True
        return False
    
    def _create_incautaciones(self, procedimiento, row_data, global_stats):
        incautaciones_raw = row_data.get('INCAUTACIONES')
        
        if incautaciones_raw and str(incautaciones_raw).strip() not in ['-', '']:
            Incautaciones.objects.create(
                procedimiento=procedimiento,
                incautaciones=self._safe_str(incautaciones_raw),
                tipo=self._safe_str(row_data.get('TIPO')),
                subtipo=self._safe_str(row_data.get('SUBTIPO')),
                cantidad=self._safe_str(row_data.get('CANTIDAD')),
                medidas=self._safe_str(row_data.get('MEDIDAS')),
                aforo=self._safe_decimal(row_data.get('AFORO')),
                observaciones=self._safe_str(row_data.get('OBSERVACIONES')),
                tipo_delito=self._safe_str(row_data.get('TIPO_DELITO')),
                juzgado_interviniente=self._safe_str(row_data.get('JUZGADO_INTERVINIENTE')),
                caratula_causa=self._safe_str(row_data.get('CARATULA_CAUSA')),
                num_causa=self._safe_str(row_data.get('NUM_CAUSA'))
            )
            global_stats['specialized_tables']['incautaciones'] += 1
            return True
        return False
    
    def _create_trata_trafico_personas(self, procedimiento, row_data, global_stats):
        TrataTraficPersonas.objects.create(
            procedimiento=procedimiento,
            tipo_delito=self._safe_str(row_data.get('TIPO_DELITO')),
            sexo_victima=self._safe_str(row_data.get('SEXO_VICTIMA')),
            genero_victima=self._safe_str(row_data.get('GENERO_VICTIMA')),
            edad_victima=self._safe_int(row_data.get('EDAD_VICTIMA')),
            nacionalidad=self._safe_str(row_data.get('NACIONALIDAD')),
            juzgado_interviniente=self._safe_str(row_data.get('JUZGADO_INTERVINIENTE')),
            caratula_causa=self._safe_str(row_data.get('CARATULA_CAUSA')),
            num_causa=self._safe_str(row_data.get('NUM_CAUSA')),
            observaciones=self._safe_str(row_data.get('OBSERVACIONES'))
        )
        global_stats['specialized_tables']['trata_personas'] += 1
        return True
    
    def _create_otros_delitos(self, procedimiento, row_data, global_stats):
        OtrosDelitos.objects.create(
            procedimiento=procedimiento,
            tipo_otro_delito=self._safe_str(row_data.get('TIPO_OTRO_DELITO')),
            sexo_victima=self._safe_str(row_data.get('SEXO_VICTIMA') or row_data.get('SEXO')),
            genero_victima=self._safe_str(row_data.get('GENERO_VICTIMA')),
            edad_victima=self._safe_int(row_data.get('EDAD_VICTIMA')),
            nacionalidad=self._safe_str(row_data.get('NACIONALIDAD')),
            observaciones=self._safe_str(row_data.get('OBSERVACIONES')),
            juzgado_interviniente=self._safe_str(row_data.get('JUZGADO_INTERVINIENTE')),
            caratula_causa=self._safe_str(row_data.get('CARATULA_CAUSA')),
            num_causa=self._safe_str(row_data.get('NUM_CAUSA'))
        )
        global_stats['specialized_tables']['otros_delitos'] += 1
        return True
    
    def _create_otros_eventos(self, procedimiento, row_data, global_stats):
        OtrosEventos.objects.create(
            procedimiento=procedimiento,
            tipo_siniestro=self._safe_str(row_data.get('TIPO_SINIESTRO')),
            cant_ilesos=self._safe_int(row_data.get('CANT_ILESOS')),
            cant_lesionados=self._safe_int(row_data.get('CANT_LESIONADOS')),
            cant_muertos=self._safe_int(row_data.get('CANT_MUERTOS')),
            observaciones=self._safe_str(row_data.get('OBSERVACIONES')),
            juzgado_interviniente=self._safe_str(row_data.get('JUZGADO_INTERVINIENTE')),
            caratula_causa=self._safe_str(row_data.get('CARATULA_CAUSA')),
            num_causa=self._safe_str(row_data.get('NUM_CAUSA'))
        )
        global_stats['specialized_tables']['otros_eventos'] += 1
        return True
    
    def _create_fallecidos(self, procedimiento, row_data, global_stats):
        Fallecidos.objects.create(
            procedimiento=procedimiento,
            servicio=self._safe_str(row_data.get('SERVICIO')),
            fuerza_de_seguridad=self._safe_str(row_data.get('FUERZA DE SEGURIDAD')),
            cant_lesionados=self._safe_int(row_data.get('CANT_LESIONADOS')),
            cant_fallecidos=self._safe_int(row_data.get('CANT_FALLECIDOS')),
            provincia_evento=self._safe_str(row_data.get('PROVINCIA')),
            departamento_evento=self._safe_str(row_data.get('DEPARTAMENTO O PARTIDO')),
            localidad_evento=self._safe_str(row_data.get('LOCALIDAD')),
            latitud_evento=self._safe_decimal(row_data.get('LATITUD')),
            longitud_evento=self._safe_decimal(row_data.get('LONGITUD')),
            fecha_evento=self._safe_str(row_data.get('FECHA')),
            hora_evento=self._safe_str(row_data.get('HORA'))
        )
        global_stats['specialized_tables']['fallecidos'] += 1
        return True
    
    def _create_abatidos(self, procedimiento, row_data, global_stats):
        Abatidos.objects.create(
            procedimiento=procedimiento,
            servicio=self._safe_str(row_data.get('SERVICIO')),
            fuerza_de_seguridad=self._safe_str(row_data.get('FUERZA DE SEGURIDAD')),
            edad=self._safe_int(row_data.get('EDAD')),
            sexo=self._safe_str(row_data.get('SEXO')),
            nacionalidad=self._safe_str(row_data.get('NACIONALIDAD')),
            terceros_damnificados=self._safe_str(row_data.get('TERCEROS DAMNIFICADOS')),
            provincia_evento=self._safe_str(row_data.get('PROVINCIA')),
            departamento_evento=self._safe_str(row_data.get('DEPARTAMENTO O PARTIDO')),
            localidad_evento=self._safe_str(row_data.get('LOCALIDAD')),
            latitud_evento=self._safe_decimal(row_data.get('LATITUD')),
            longitud_evento=self._safe_decimal(row_data.get('LONGITUD')),
            fecha_evento=self._safe_str(row_data.get('FECHA')),
            hora_evento=self._safe_str(row_data.get('HORA'))
        )
        global_stats['specialized_tables']['abatidos'] += 1
        return True
    
    def _create_codigo_operativo(self, procedimiento, row_data, global_stats):
        if not hasattr(procedimiento, 'codigo_operativo'):
            CodigoOperativo.objects.create(
                procedimiento=procedimiento,
                codigo_operativo=self._safe_str(row_data.get('CODIGO_OPERATIVO'))
            )
            global_stats['specialized_tables']['codigos_operativos'] += 1
        return True
    
    def _safe_str(self, value):
        if value is None or str(value).strip() in ['None', 'null', '-', '']:
            return None
        return str(value).strip()

    def _safe_int(self, value):
        if not value or str(value).strip() in ['-', '', 'None', 'null']:
            return None
        try:
            return int(float(str(value)))
        except (ValueError, TypeError):
            return None

    def _safe_decimal(self, value):
        if not value or str(value).strip() in ['-', '', 'None', 'null']:
            return None
        try:
            clean_value = str(value).replace(',', '.').strip()
            return Decimal(clean_value)
        except (InvalidOperation, ValueError):
            return None
    
    def _generate_filtering_statistics(self, stats):
        specialized_tables = stats.get('specialized_tables', {})
        total_geografia = specialized_tables.get('geografia_procedimientos', 0)
        
        criterios = {
            'incautaciones': 'INCAUTACIONES ≠ "-"',
            'detenidos': 'EDAD ≠ "-"', 
            'controlados': 'VEHICULOS_CONTROLADOS ≠ "-" OR PERSONAS_CONTROLADAS ≠ "-"',
            'afectados': 'CANT_EFECTIVOS > 0',
            'trata_personas': 'Todos los registros (sin filtrado)',
            'otros_delitos': 'Todos los registros (sin filtrado)',
            'otros_eventos': 'Todos los registros (sin filtrado)',
            'fallecidos': 'Todos los registros (sin filtrado)',
            'abatidos': 'Todos los registros (sin filtrado)',
            'codigos_operativos': 'Todos los registros (sin filtrado)'
        }
        
        filtering_details = []
        filtered_tables = ['incautaciones', 'vehiculos_controladas', 'personal_afectados', 'detenidos']
        
        for table_key, count in specialized_tables.items():
            if table_key == 'geografia_procedimientos':
                continue
                
            table_display_names = {
                'vehiculos_controladas': 'controlados',
                'personal_afectados': 'afectados',
                'detenidos': 'detenidos',
                'incautaciones': 'incautaciones',
                'trata_personas': 'trata_personas',
                'otros_delitos': 'otros_delitos',
                'otros_eventos': 'otros_eventos',
                'fallecidos': 'fallecidos',
                'abatidos': 'abatidos',
                'codigos_operativos': 'codigos_operativos'
            }
            
            display_name = table_display_names.get(table_key, table_key)
            
            if table_key in filtered_tables:
                registros_omitidos = max(0, total_geografia - count)
                porcentaje_reduccion = (registros_omitidos / total_geografia * 100) if total_geografia > 0 else 0
            else:
                registros_omitidos = 0
                porcentaje_reduccion = 0
            
            filtering_details.append({
                'tabla': display_name.replace('_', ' ').title(),
                'registros_total': total_geografia,
                'registros_filtrados': count,
                'registros_omitidos': registros_omitidos,
                'porcentaje_reduccion': round(porcentaje_reduccion, 2),
                'criterio_filtrado': criterios.get(display_name, 'Sin criterio especificado')
            })
        
        return filtering_details
